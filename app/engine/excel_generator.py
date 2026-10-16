import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.schemas import SignalModel

def build_scenario_input_workbook(scenario_id: str, signals: List[SignalModel], scenario_title: str, scenario_description: str) -> io.BytesIO:
    """
    Builds a professional, formatted Excel workbook for sample signal input data.
    Includes a dedicated README worksheet explaining the scenario to judges/evaluators,
    and a formatted 'Signal Register' sheet with frozen headers and auto-sized columns.
    """
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------------------
    # Sheet 1: README & Scenario Guide
    # -------------------------------------------------------------------------
    ws_readme = wb.active
    ws_readme.title = "Scenario Guide"
    ws_readme.views.sheetView[0].showGridLines = True

    # Title styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1F2423")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="247C70")
    body_font = Font(name="Calibri", size=10, color="333333")
    bold_body = Font(name="Calibri", size=10, bold=True, color="1F2423")
    box_fill = PatternFill(start_color="F4F6F4", end_color="F4F6F4", fill_type="solid")
    accent_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='C8CECA'),
        right=Side(style='thin', color='C8CECA'),
        top=Side(style='thin', color='C8CECA'),
        bottom=Side(style='thin', color='C8CECA')
    )

    ws_readme["B2"] = "IO SIGNAL SUDOKU — Industrial I/O Allocation Platform"
    ws_readme["B2"].font = title_font
    ws_readme["B3"] = f"DEMO SCENARIO DATASET: {scenario_title}"
    ws_readme["B3"].font = subtitle_font

    ws_readme["B5"] = "SCENARIO OVERVIEW"
    ws_readme["B5"].font = bold_body
    ws_readme["B6"] = scenario_description
    ws_readme["B6"].font = body_font

    # Info Box
    info_rows = [
        ("Platform Target", "Emerson DeltaV Electronic Marshaling Cabinet CAB-01 (64 CHARMs)"),
        ("Optimization Engine", "Google OR-Tools CP-SAT v9.10+ (Constraint Programming)"),
        ("Independent Validator", "IEC 61131-2, IEC 60079-11 (Ex-ia), IEEE 518 Noise Segregation"),
        ("Data Origin", "SYNTHETIC DEMO BENCHMARK — Clearly labeled for hackathon evaluation"),
        ("Signal Count", f"{len(signals)} instrumentation loops in this register")
    ]

    start_r = 8
    for label, val in info_rows:
        ws_readme.cell(row=start_r, column=2, value=label).font = bold_body
        ws_readme.cell(row=start_r, column=2).fill = box_fill
        ws_readme.cell(row=start_r, column=2).border = thin_border
        
        ws_readme.cell(row=start_r, column=3, value=val).font = body_font
        ws_readme.cell(row=start_r, column=3).fill = box_fill
        ws_readme.cell(row=start_r, column=3).border = thin_border
        start_r += 1

    ws_readme.column_dimensions["B"].width = 24
    ws_readme.column_dimensions["C"].width = 65

    # -------------------------------------------------------------------------
    # Sheet 2: Signal Register (Importable Data)
    # -------------------------------------------------------------------------
    ws_data = wb.create_sheet(title="Signal Register")
    ws_data.views.sheetView[0].showGridLines = True

    headers = [
        "tag",
        "description",
        "signal_type",
        "voltage",
        "is_intrinsic_safety",
        "safety_integrity_level",
        "multicore_cable_id",
        "cable_pair_index",
        "power_dissipation_watts",
        "locked",
        "assigned_position_id"
    ]

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="247C70", end_color="247C70", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAF9", end_color="F8FAF9", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Write headers
    for col_idx, h in enumerate(headers, start=1):
        c = ws_data.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    # Write data rows
    for r_idx, s in enumerate(signals, start=2):
        row_values = [
            s.tag,
            s.description,
            s.signal_type,
            s.voltage or "24VDC",
            "TRUE" if s.is_intrinsic_safety else "FALSE",
            s.safety_integrity_level or "BPCS",
            s.multicore_cable_id or "",
            s.cable_pair_index or "",
            s.power_dissipation_watts,
            "TRUE" if s.locked else "FALSE",
            s.assigned_position_id or ""
        ]
        
        fill = alt_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
        
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws_data.cell(row=r_idx, column=col_idx, value=val)
            cell.font = body_font
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
            if col_idx in [1, 3, 4, 5, 6, 7, 8, 10, 11]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Freeze header
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(signals) + 1}"

    # Auto-adjust column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def get_sample_workbook_by_name(filename: str) -> io.BytesIO:
    """
    Returns an Excel BytesIO stream for the 5 official scenario sample workbooks:
    - Scenario_A_Clean_Input.xlsx
    - Scenario_B_Warnings_Input.xlsx
    - Scenario_C_Conflict_Input.xlsx
    - Scenario_D_Revision_Baseline.xlsx
    - Scenario_D_Revision_Change.xlsx
    """
    from app.engine.demo_data import (
        get_scenario_a_signals,
        get_scenario_b_warning_signals,
        get_scenario_c_conflict_signals,
        get_scenario_d_signals
    )

    clean_name = filename.strip()
    if clean_name in ("Scenario_A_Clean_Input.xlsx", "SCENARIO_A"):
        return build_scenario_input_workbook(
            scenario_id="SCENARIO_A",
            signals=get_scenario_a_signals(),
            scenario_title="Scenario A: Clean Allocation Baseline",
            scenario_description="Standard baseline instrumentation register (45 loops). Demonstrates optimal allocation with 0 violations and 100% feasibility under Google OR-Tools CP-SAT."
        )
    elif clean_name in ("Scenario_B_Warnings_Input.xlsx", "SCENARIO_B"):
        return build_scenario_input_workbook(
            scenario_id="SCENARIO_B",
            signals=get_scenario_b_warning_signals(),
            scenario_title="Scenario B: Advisory Warnings (0 Violations)",
            scenario_description="Advisory warnings fixture (45 loops). Satisfies all mandatory constraints (100% valid), but triggers genuine RULE-SOFT-02 (multicore cable pair split) and RULE-HARD-05 (expansion reserve slot utilized)."
        )
    elif clean_name in ("Scenario_C_Conflict_Input.xlsx", "SCENARIO_C"):
        return build_scenario_input_workbook(
            scenario_id="SCENARIO_C",
            signals=get_scenario_c_conflict_signals(),
            scenario_title="Scenario C: Engineering Constraint Conflict",
            scenario_description="Safety conflict fixture (46 loops). Hazardous Zone 0 Ex-ia transmitter PT-FLARE-IS is locked to standard non-IS baseplate BP-01, triggering CP-SAT infeasibility and IEC 60079-11 physical barrier violation."
        )
    elif clean_name in ("Scenario_D_Revision_Baseline.xlsx", "SCENARIO_D_BASELINE"):
        return build_scenario_input_workbook(
            scenario_id="SCENARIO_D",
            signals=get_scenario_a_signals(),
            scenario_title="Scenario D: Revision Baseline (Rev A)",
            scenario_description="Approved baseline allocation (45 loops) for DeltaV cabinet CAB-01 prior to engineering change order ECO-482."
        )
    elif clean_name in ("Scenario_D_Revision_Change.xlsx", "SCENARIO_D", "SCENARIO_D_CHANGE"):
        return build_scenario_input_workbook(
            scenario_id="SCENARIO_D",
            signals=get_scenario_d_signals(),
            scenario_title="Scenario D: Revision Change Order (Rev B ECO)",
            scenario_description="Revised instrumentation register (49 loops: +5 additions, -1 deletion, PT-102 upgraded to Ex-ia). Minimal-Repair planner minimizes field rewire churn."
        )
    else:
        raise ValueError(f"Unknown sample workbook filename: {filename}")

def generate_all_sample_files(output_dir: str = "app/static/samples") -> Dict[str, str]:
    """Generates and writes all 5 sample Excel workbooks to static storage."""
    from pathlib import Path
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    filenames = [
        "Scenario_A_Clean_Input.xlsx",
        "Scenario_B_Warnings_Input.xlsx",
        "Scenario_C_Conflict_Input.xlsx",
        "Scenario_D_Revision_Baseline.xlsx",
        "Scenario_D_Revision_Change.xlsx"
    ]
    created = {}
    for fn in filenames:
        buf = get_sample_workbook_by_name(fn)
        file_dest = out_path / fn
        with open(file_dest, "wb") as f:
            f.write(buf.getvalue())
        created[fn] = str(file_dest)
    return created
